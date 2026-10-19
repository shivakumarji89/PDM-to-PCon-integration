"""File/stream output chunk of ExportThread.cs, ported as ``IoMixin``.

C# ranges ported (ExportThread.cs): removeNonStandardCharacters (293-331), writeToStream (333-355),
allocateFiles (437-466), WriteToSheetSlow (468-485), WriteToSheet (487-521), ConvertToXLS (523-707),
writeData (752-787), logMessage (789-812), writePriceFormulaData (986-1036),
writeProductCodeData (1038-1088).

Writers
-------
Legacy ``StreamWriter(path, append: false)`` (UTF-8, no BOM) becomes
``open(path, 'w', encoding='utf-8', newline='')`` (``open_writer``); ``newline=''`` keeps the
literal ``\\r\\n`` the legacy strings contain.

DELIBERATE DEVIATION: Excel COM interop -> openpyxl
---------------------------------------------------
``ConvertToXLS`` drives Excel through COM and saves ``<exportDirectory><exportFileName>.xls``.
The port has no Excel dependency: it writes ``<exportDirectory><exportFileName>.xlsx`` with
openpyxl (output file name decision: same stem as the C# ``.xls`` name, extension ``.xlsx``,
because an xlsx payload under an ``.xls`` name makes Excel warn). Data semantics kept:

* sheet names, which sheets exist per mode, and their final order. ``Worksheets.Add()`` with no
  position inserts before the active sheet, so the *last added sheet ends up first*:
    - syteline, not CSI:  Items, Materials, SPextract, Planning Data, Descriptions
    - syteline + CSI:     Items, Materials
    - not syteline:       Items, Materials, PriceFormula, ProductCode, SPextract
  (the initial "Sheet1" is created and deleted exactly as legacy, so it never survives);
* Items/Materials come from ``ReadToEnd().Split('\\n')`` (so a trailing "\\r" stays on the last cell of
  every row and the file's final empty segment becomes a last, empty row); the other sheets come from
  ``File.ReadAllLines`` (no "\\r", no phantom last row);
* the column count is that of the first line; longer rows are truncated to it, shorter rows leave the
  remaining cells empty (``WriteToSheet``);
* Materials column index 3 (4th column, header row included) is text-forced. Excel does this with a
  leading apostrophe (a prefix marker, not cell content); here the value is stored as a string with
  ``quotePrefix`` so the stored content is the same;
* AutoFit is approximated by max text length; the A1 selection of Items is the active sheet.

Not reproduced: Excel's implicit typing of other cells (numeric-looking strings, dates and "=..."
text become numbers/dates/formulas under COM; here every cell stays a string).
A failure in conversion is recorded via ``report_swallowed`` (legacy shows a MsgBox and still
reports completion) and ALSO appended to ``self.warnings`` so the UI can show partial success.
"""
from __future__ import annotations

import os
import re
from typing import IO

from services.syteline_export.pipeline.netsem import list_contains, substring, tostr

_STANDARD_CHARACTERS = (
    "0123456789abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ"
    "-+=_!\"£$%^&*()[]<>,./?#~{}:;@'`|\\ "
)
_LINE_SPLIT = re.compile(r"\r\n|\r|\n")


def _read_all_text(path: str) -> str:
    """StreamReader.ReadToEnd(): UTF-8 (BOM tolerated), no newline translation."""
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        return handle.read()


def _read_all_lines(path: str) -> list[str]:
    """File.ReadAllLines(): splits on \\r\\n, \\r, \\n; no trailing empty element."""
    lines = _LINE_SPLIT.split(_read_all_text(path))
    if lines and lines[-1] == "":
        lines.pop()
    return lines


_XLSX_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")
_MAX_SQL_PARAMETERS = 2000  # SQL Server allows 2100 parameters per statement; the legacy inlined literals


class IoMixin:
    # ExportThread.cs:293-331
    def remove_non_standard_characters(self, data: str) -> str:
        text: list[str] = []
        try:
            text2 = _STANDARD_CHARACTERS
            for i in range(0, len(data)):
                ch = data[i]
                if text2.find(ch) > -1:
                    text.append(ch)
                elif ch == "\r":
                    text.append(ch)
                elif ch == "\n":
                    text.append(ch)
                elif ch == "\t":
                    text.append(ch)
        except Exception as exc:  # legacy MsgBox(ex.ToString()); returns what was built so far
            self.report_swallowed("removeNonStandardCharacters", exc)
        return "".join(text)

    # ExportThread.cs:333-355
    def write_to_stream(self, sw: IO[str], data: str, debugcontext: str) -> None:
        try:
            data = self.remove_non_standard_characters(data)
            data = data.replace("\r\n\r\n", "\r\n")
            # debug-only block omitted (flag default false): SytelineExport.debugDuplicateLines
            sw.write(data)
        except Exception as exc:
            self.report_swallowed("writeToStream", exc)

    # StreamWriter(path, append: false) with legacy encoding (UTF-8, no BOM)
    def open_writer(self, path: str) -> IO[str]:
        return open(path, "w", encoding="utf-8", errors="replace", newline="")

    # ExportThread.cs:437-466
    def allocate_files(self) -> None:
        if self._millerCAD:
            self._itemFile = self._exportDirectory + "b-" + self._site.lower() + self._currency + ".asc"
            self._materialFile = self._exportDirectory + "o-" + self._site.lower() + self._currency + ".asc"
            if self._exportFileName != "b-" + self._site.lower() + self._currency:
                self._itemFile = (
                    self._itemFile.replace(".asc", "") + "_"
                    + self._exportFileName.replace("b-" + self._site.lower() + self._currency, "").lstrip("_")
                    + ".asc"
                )
                self._materialFile = (
                    self._materialFile.replace(".asc", "") + "_"
                    + self._exportFileName.replace("b-" + self._site.lower() + self._currency, "").lstrip("_")
                    + ".asc"
                )
            return
        text = ""
        if self._formatCSI:
            text = "_csi"
        if not os.path.isdir(self._exportDirectory + "temp\\"):
            os.makedirs(self._exportDirectory + "temp\\", exist_ok=True)
        self._itemFile = self._exportDirectory + "temp\\" + self._exportFileName + "_item_data" + text + ".txt"
        self._materialFile = self._exportDirectory + "temp\\" + self._exportFileName + "_material_data" + text + ".txt"
        self._superprodFile = self._exportDirectory + "temp\\" + self._exportFileName + "_superprod_data" + text + ".txt"
        self._descriptionFile = self._exportDirectory + "temp\\" + self._exportFileName + "_description_data" + text + ".txt"
        self._planningFile = self._exportDirectory + "temp\\" + self._exportFileName + "_planning_data" + text + ".txt"
        self._priceformulaFile = self._exportDirectory + "temp\\" + self._exportFileName + "_priceformula_data" + text + ".txt"
        self._productcodeFile = self._exportDirectory + "temp\\" + self._exportFileName + "_productcode_data" + text + ".txt"

    # ExportThread.cs:468-485 (no caller in the legacy file; ported for completeness)
    def write_to_sheet_slow(self, sheet, data_lines: list[str], context: str) -> None:
        num = 1
        for text in data_lines:
            self.update_status_label(
                "writing row " + str(num) + " of " + str(len(data_lines)) + " (" + context + ")"
            )
            array = text.replace("\r", "").split("\t")
            for j in range(0, len(array)):
                sheet.cell(row=num, column=j + 1).value = array[j]
            num += 1

    # ExportThread.cs:487-521
    def write_to_sheet(self, sheet, data_lines: list[str], context: str) -> None:
        num = len(data_lines)  # row count
        if num == 0:
            return
        array = data_lines[0].split("\t")
        num2 = len(array)  # column count taken from the first line
        for i in range(0, num):
            array3 = data_lines[i].split("\t")
            num4 = min(len(array3), num2) - 1
            for j in range(0, num4 + 1):
                cell = sheet.cell(row=i + 1, column=j + 1)
                cell.value = _XLSX_ILLEGAL.sub("", array3[j])  # openpyxl rejects control characters Excel accepts
                cell.data_type = "s"  # never a formula / number (deviation: no Excel auto-typing)
                if context == "Materials" and j == 3:
                    cell.quotePrefix = True  # legacy: "'" + value (text forcing)

    @staticmethod
    def _autofit(sheet) -> None:
        """Worksheet.Columns.AutoFit approximation."""
        widths: dict[int, int] = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
        for column, width in widths.items():
            sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = min(width + 2, 255)

    # ExportThread.cs:523-707 -> openpyxl (see module docstring)
    def convert_to_xls(self) -> None:
        workbook = None
        try:
            from openpyxl import Workbook  # deviation: no Excel COM

            text = self._exportDirectory + self._exportFileName + ".xlsx"  # legacy: ".xls"
            text2 = _read_all_text(self._itemFile)
            text3 = _read_all_text(self._materialFile)
            workbook = Workbook()  # SheetsInNewWorkbook = 1
            workbook.active.title = "Sheet1"
            worksheet = worksheet2 = worksheet3 = worksheet4 = worksheet5 = None

            def add(name: str):
                # Worksheets.Add() with no position inserts before the active (= previously added) sheet
                return workbook.create_sheet(name, 0)

            if self._syteline:
                if not self._formatCSI:
                    dataLines = _read_all_lines(self._superprodFile)
                    dataLines2 = _read_all_lines(self._planningFile)
                    dataLines3 = _read_all_lines(self._descriptionFile)
                    worksheet5 = add("Descriptions")
                    self.write_to_sheet(worksheet5, dataLines3, "Descriptions")
                    worksheet4 = add("Planning Data")
                    self.write_to_sheet(worksheet4, dataLines2, "Planning Data")
                    worksheet3 = add("SPextract")
                    self.write_to_sheet(worksheet3, dataLines, "SPextract")
            else:
                dataLines4 = _read_all_lines(self._superprodFile)
                dataLines5 = _read_all_lines(self._productcodeFile)
                dataLines6 = _read_all_lines(self._priceformulaFile)
                worksheet5 = add("SPextract")
                self.write_to_sheet(worksheet5, dataLines4, "SPextract")
                worksheet4 = add("ProductCode")
                self.write_to_sheet(worksheet4, dataLines5, "ProductCode")
                worksheet3 = add("PriceFormula")
                self.write_to_sheet(worksheet3, dataLines6, "PriceFormula")
            dataLines7 = text3.split("\n")
            dataLines8 = text2.split("\n")
            worksheet2 = add("Materials")
            self.write_to_sheet(worksheet2, dataLines7, "Materials")
            worksheet = add("Items")
            self.write_to_sheet(worksheet, dataLines8, "Items")
            self._autofit(worksheet)
            self._autofit(worksheet2)
            if not self._formatCSI:
                self._autofit(worksheet3)
                self._autofit(worksheet4)
                if not self._syteline:
                    self._autofit(worksheet5)
            for worksheet6 in list(workbook.worksheets):
                if worksheet6.title == "Sheet1":
                    workbook.remove(worksheet6)
                    break
            for sheet in workbook.worksheets:
                sheet.sheet_view.tabSelected = False
            workbook.active = workbook.worksheets.index(worksheet)  # Items!A1 selected
            worksheet.sheet_view.tabSelected = True
            if os.path.exists(text):
                self.update_status_label("deleting prev instance")
                os.remove(text)
            self.update_status_label("saving file as xlsx")
            workbook.save(text)
        except Exception as exc:
            self.report_swallowed("ERROR during XLS export", exc)
            self.warnings.append(
                "XLS/XLSX conversion failed (text export files were still written): "
                + type(exc).__name__ + ": " + str(exc)
            )
        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:  # legacy swallows release errors silently
                    pass

    # ExportThread.cs:752-787
    def write_data(self, mywriter: IO[str], mydata: str, isItem: bool) -> None:
        try:
            if mydata.startswith("9-F"):
                text = substring(mydata, 0, mydata.find("\t"))
                if list_contains(self._materialsOnItemTab, text):
                    return
                self._materialsOnItemTab.append(text)
            if isItem:
                if list_contains(self._itemsWritten, mydata):
                    return
                self._itemsWritten.append(mydata)
            if (mydata.find("BOCFAB-OF") > -1) & (mydata.find("2-HAR") > -1):
                mydata = mydata.replace("-BOCFAB-", "-")
            self.write_to_stream(mywriter, mydata.replace("HAR-F", "H-F"), "aaaaaa")
            mywriter.flush()
        except Exception as exc:
            self.report_swallowed("writeData", exc)

    # ExportThread.cs:789-812. Legacy writes only when SytelineExport.logOutput (default false, never
    # set) and Environment.UserName == "dbacw8" (developer-machine hack, C:\Downloads\log.txt): omitted.
    @staticmethod
    def log_message(msg: str, append: bool = True) -> None:
        return

    # ExportThread.cs:986-1036
    def write_price_formula_data(self) -> None:
        # legacy opens the writer before the try: a failure here propagates to the caller
        streamWriter = self.open_writer(self._priceformulaFile)
        try:
            streamWriter.write("Site\tCustPriceCode\tItemPriceCode\tPriceFormula\tPriceFactor\tEffectiveDate\r\n")
            if len(self._priceCodes) > _MAX_SQL_PARAMETERS:
                raise ValueError(f"Too many price codes ({len(self._priceCodes)}) for one bound-parameter query.")
            if len(self._priceCodes) > 0:
                text = (
                    "SELECT Site.Site, pm.CustPriceCode, pm.ItemPriceCode, pf.PriceFormula, pf.FirstPrice, "
                    "CONVERT(VARCHAR(11), pf.EffectiveDate, 106) AS EffDate FROM PriceFormula pf WITH (NOLOCK) "
                    "INNER JOIN PriceMatrix pm ON pf.PriceFormula = pm.PriceFormula "
                    "INNER JOIN Site ON pf.SiteId = Site.SiteId WHERE Site.SiteId = ? AND ("
                )
                params: list = [self._siteId]
                for i in range(0, len(self._priceCodes)):
                    text = text + "pm.ItemPriceCode = ?"
                    params.append(tostr(self._priceCodes[i]))
                    if i < len(self._priceCodes) - 1:
                        text += " OR "
                text += ") ORDER BY pm.CustPriceCode, pf.EffectiveDate DESC"
                # legacy used a fresh ConnectionFactory connection (not one of the five run connections)
                rows = self.q("c_priceformula", text, tuple(params))
                for row in rows:
                    streamWriter.write(
                        tostr(row["Site"]) + "\t" + tostr(row["CustPriceCode"]) + "\t"
                        + tostr(row["ItemPriceCode"]) + "\t" + tostr(row["PriceFormula"]) + "\t"
                        + tostr(row["FirstPrice"]) + "\t" + tostr(row["EffDate"]) + "\r\n"
                    )
        except Exception as exc:
            self.report_swallowed("writePriceFormulaData", exc)
        finally:
            streamWriter.close()

    # ExportThread.cs:1038-1088
    def write_product_code_data(self) -> None:
        streamWriter = self.open_writer(self._productcodeFile)
        try:
            # LEGACY DEFECT kept: 5 header columns (UnitCode) but rows carry only 4 values
            streamWriter.write("Site\tProductCode\tDescription\tPriceCode\tUnitCode\r\n")
            if len(self._productCodes) > _MAX_SQL_PARAMETERS:
                raise ValueError(f"Too many product codes ({len(self._productCodes)}) for one bound-parameter query.")
            if len(self._productCodes) > 0:
                text = (
                    "SELECT Site.Site, pc.Product_Code, pc.Description, pc.PriceCode FROM Product_Code pc "
                    "WITH (NOLOCK) INNER JOIN Site ON pc.SiteId = Site.SiteId WHERE Site.SiteId = ? AND ("
                )
                params: list = [self._siteId]
                for i in range(0, len(self._productCodes)):
                    text = text + "pc.Product_Code = ?"
                    params.append(tostr(self._productCodes[i]))
                    if i < len(self._productCodes) - 1:
                        text += " OR "
                text += ") ORDER BY pc.Product_Code"
                rows = self.q("c_productcode", text, tuple(params))
                for row in rows:
                    streamWriter.write(
                        tostr(row["Site"]) + "\t" + tostr(row["Product_Code"]) + "\t"
                        + tostr(row["Description"]) + "\t" + tostr(row["PriceCode"]) + "\r\n"
                    )
        except Exception as exc:
            self.report_swallowed("writeProductCodeData", exc)
        finally:
            streamWriter.close()
