"""Export a complete Snapshot to a review-friendly multi-sheet Excel workbook."""
from __future__ import annotations

import json
import re
from dataclasses import fields, is_dataclass
from enum import Enum
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def _json_value(value: Any) -> str:
    def normalise(item):
        if isinstance(item, Enum):
            return item.value
        if is_dataclass(item):
            return {f.name: normalise(getattr(item, f.name)) for f in fields(item)}
        if isinstance(item, dict):
            return {str(k): normalise(v) for k, v in item.items()}
        if isinstance(item, (list, tuple, set)):
            return [normalise(v) for v in item]
        if item is None or isinstance(item, (str, int, float, bool)):
            return item
        return str(item)

    return json.dumps(normalise(value), ensure_ascii=False, default=str)


def _flat_row(value: Any) -> dict[str, Any]:
    """Flatten nested records into Excel columns; complex cells remain JSON."""
    if is_dataclass(value):
        source = {f.name: getattr(value, f.name) for f in fields(value)}
    elif isinstance(value, dict):
        source = value
    else:
        return {"Value": value if value is None or isinstance(value, (str, int, float, bool)) else _json_value(value)}

    result = {}
    for key, item in source.items():
        name = str(key)
        if item is None or isinstance(item, (str, int, float, bool)):
            result[name] = item
        elif isinstance(item, Enum):
            result[name] = item.value
        else:
            result[name] = _json_value(item)
    return result


class SnapshotExcelExportService:
    """Write all top-level snapshot data to separate, reviewable worksheets."""

    def export(self, snapshot, destination: str | Path) -> Path:
        if snapshot is None:
            raise ValueError("No snapshot is loaded.")

        workbook = Workbook()
        index = workbook.active
        index.title = "Workbook Info"
        index.append(["Sheet", "Records", "Contents"])
        sheets = []

        def add_sheet(title: str, rows: list[dict], description: str):
            safe = re.sub(r"[^A-Za-z0-9 _-]", "_", title).strip() or "Data"
            safe = safe[:31]
            base, suffix = safe, 2
            while safe in workbook.sheetnames:
                tail = f"_{suffix}"
                safe = base[:31 - len(tail)] + tail
                suffix += 1
            ws = workbook.create_sheet(safe)
            if not rows:
                ws.append(["No records"])
            else:
                columns = []
                for row in rows:
                    for key in row:
                        if key not in columns:
                            columns.append(key)
                ws.append(columns)
                for row in rows:
                    ws.append([row.get(col) for col in columns])
            sheets.append((safe, len(rows), description))
            return ws

        # One row per scalar leaf means no Snapshot field is silently omitted.
        leaves = []
        collections = []
        def walk(value, path):
            if is_dataclass(value):
                for field_info in fields(value):
                    walk(getattr(value, field_info.name), f"{path}.{field_info.name}".strip("."))
            elif isinstance(value, dict):
                if not value:
                    leaves.append({"Field Path": path, "Value": "{}"})
                for key, child in value.items():
                    walk(child, f"{path}[{key}]")
            elif isinstance(value, (list, tuple)):
                if not value:
                    leaves.append({"Field Path": path, "Value": "[]"})
                else:
                    rows = []
                    for idx, item in enumerate(value):
                        row = _flat_row(item)
                        row = {"Index": idx, **row}
                        rows.append(row)
                        # Recurse so nested lists/maps are also individually reviewable.
                        if is_dataclass(item) or isinstance(item, (dict, list, tuple)):
                            walk(item, f"{path}[{idx}]")
                    collections.append((path, rows))
            elif isinstance(value, Enum):
                leaves.append({"Field Path": path, "Value": str(value.value)})
            elif value is None or isinstance(value, (str, int, float, bool)):
                leaves.append({"Field Path": path, "Value": value})
            else:
                leaves.append({"Field Path": path, "Value": str(value)})

        walk(snapshot, "Snapshot")
        add_sheet("All Fields", leaves, "Complete field-path/value inventory, including nested structures.")
        for path, rows in collections:
            # Sheet names are uniquified by add_sheet; collection name is kept in the index.
            add_sheet(path.replace("Snapshot.", ""), rows, f"Collection from {path}")

        for name, count, description in sheets:
            index.append([name, count, description])

        for ws in workbook.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions if ws.max_row > 1 else "A1"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="355070")
            for col_cells in ws.columns:
                letter = get_column_letter(col_cells[0].column)
                max_width = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells[:200])
                ws.column_dimensions[letter].width = min(max(max_width + 2, 12), 48)

        target = Path(destination)
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target)
        return target
