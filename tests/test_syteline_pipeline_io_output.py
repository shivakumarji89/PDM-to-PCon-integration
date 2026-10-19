"""Tests for pipeline/io_output.py (ExportThread.cs 293-355, 437-521, 523-707, 752-812, 986-1088).

All fixtures are synthetic. Expected values derive from the C# (line numbers cited).
"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from services.syteline_export.pipeline.db import ScriptedExportDb
from services.syteline_export.pipeline.io_output import IoMixin
from services.syteline_export.pipeline.netsem import LegacyRangeError
from services.syteline_export.pipeline.run_base import ExportRunBase


class H(IoMixin, ExportRunBase):
    pass


def make(tmp_path, script=None, **fields):
    h = H(ScriptedExportDb(script or {}))
    h._exportDirectory = str(tmp_path) + "\\"
    for k, v in fields.items():
        setattr(h, k, v)
    return h


def read(path):
    with open(path, "r", encoding="utf-8", newline="") as f:
        return f.read()


# --- removeNonStandardCharacters (293-331) -------------------------------------------------------
def test_remove_non_standard_characters(tmp_path):
    h = make(tmp_path)
    # keeps listed chars incl. pound, backslash, space, CR LF TAB; drops others (e, acute, euro, NUL)
    assert h.remove_non_standard_characters("aé€\x00b£\\ \r\n\t-") == "ab£\\ \r\n\t-"
    assert h.errors == []


def test_remove_non_standard_characters_exception_is_swallowed(tmp_path):
    h = make(tmp_path)
    assert h.remove_non_standard_characters(None) == ""  # len(None) -> caught (322-328)
    assert len(h.errors) == 1


# --- writeToStream (333-355) ---------------------------------------------------------------------
def test_write_to_stream_collapses_double_crlf_once(tmp_path):
    h = make(tmp_path)
    p = tmp_path / "a.txt"
    with h.open_writer(str(p)) as f:
        h.write_to_stream(f, "a\r\n\r\nb\r\n\r\n\r\n\r\nc€", "x")
    # "\r\n\r\n\r\n\r\n" -> Replace is non-recursive -> "\r\n\r\n"
    assert read(p) == "a\r\nb\r\n\r\nc"


def test_write_to_stream_error_swallowed(tmp_path):
    h = make(tmp_path)

    class Bad:
        def write(self, s):
            raise OSError("disk")

    h.write_to_stream(Bad(), "x", "ctx")
    assert len(h.errors) == 1


def test_writer_is_utf8_without_bom_and_keeps_crlf(tmp_path):
    h = make(tmp_path)
    p = tmp_path / "b.txt"
    with h.open_writer(str(p)) as f:
        f.write("£\r\n")
    assert p.read_bytes() == "£\r\n".encode("utf-8")


# --- allocateFiles (437-466) ---------------------------------------------------------------------
def test_allocate_files_normal_creates_temp(tmp_path):
    h = make(tmp_path, _exportFileName="PDMExport")
    h.allocate_files()
    d = str(tmp_path) + "\\temp\\"
    assert (tmp_path / "temp").is_dir()
    assert h._itemFile == d + "PDMExport_item_data.txt"
    assert h._materialFile == d + "PDMExport_material_data.txt"
    assert h._superprodFile == d + "PDMExport_superprod_data.txt"
    assert h._descriptionFile == d + "PDMExport_description_data.txt"
    assert h._planningFile == d + "PDMExport_planning_data.txt"
    assert h._priceformulaFile == d + "PDMExport_priceformula_data.txt"
    assert h._productcodeFile == d + "PDMExport_productcode_data.txt"


def test_allocate_files_csi_suffix(tmp_path):
    h = make(tmp_path, _formatCSI=True)
    h.allocate_files()
    assert h._itemFile.endswith("PDMExport_item_data_csi.txt")
    assert h._productcodeFile.endswith("PDMExport_productcode_data_csi.txt")


def test_allocate_files_millercad_default_name(tmp_path):
    h = make(tmp_path, _millerCAD=True, _site="UK", _currency="GBP", _exportFileName="b-ukGBP")
    h.allocate_files()
    assert h._itemFile == str(tmp_path) + "\\b-ukGBP.asc"      # 441
    assert h._materialFile == str(tmp_path) + "\\o-ukGBP.asc"  # 442
    assert not (tmp_path / "temp").exists()                    # millerCAD returns before CreateDirectory


def test_allocate_files_millercad_custom_name(tmp_path):
    h = make(tmp_path, _millerCAD=True, _site="UK", _currency="GBP", _exportFileName="b-ukGBP__spring")
    h.allocate_files()
    # 445/446: Replace removes the prefix, TrimStart('_') strips leading underscores
    assert h._itemFile == str(tmp_path) + "\\b-ukGBP_spring.asc"
    assert h._materialFile == str(tmp_path) + "\\o-ukGBP_spring.asc"
    h2 = make(tmp_path, _millerCAD=True, _site="UK", _currency="GBP", _exportFileName="other")
    h2.allocate_files()
    assert h2._itemFile == str(tmp_path) + "\\b-ukGBP_other.asc"


# --- writeData (752-787) -------------------------------------------------------------------------
def test_write_data_item_dedup_and_replacements(tmp_path):
    h = make(tmp_path)
    p = tmp_path / "w.txt"
    with h.open_writer(str(p)) as f:
        h.write_data(f, "X-HAR-F1\t1\r\n", True)
        h.write_data(f, "X-HAR-F1\t1\r\n", True)  # dup item skipped (767)
        h.write_data(f, "X-HAR-F1\t1\r\n", False)  # not an item -> written again
    assert read(p) == "X-H-F1\t1\r\n" * 2  # HAR-F -> H-F (777)


def test_write_data_bocfab_rule(tmp_path):
    h = make(tmp_path)
    p = tmp_path / "w.txt"
    with h.open_writer(str(p)) as f:
        h.write_data(f, "2-HAR-BOCFAB-OF\r\n", False)  # both substrings present (773-776)
        h.write_data(f, "1-BOCFAB-OF\r\n", False)      # no 2-HAR: untouched
    assert read(p) == "2-HAR-OF\r\n" + "1-BOCFAB-OF\r\n"


def test_write_data_9f_dedups_on_first_column(tmp_path):
    h = make(tmp_path)
    p = tmp_path / "w.txt"
    with h.open_writer(str(p)) as f:
        h.write_data(f, "9-F1\tA\r\n", False)
        h.write_data(f, "9-F1\tB\r\n", False)  # same key skipped (759)
        h.write_data(f, "9-F2\tB\r\n", False)
    assert read(p) == "9-F1\tA\r\n9-F2\tB\r\n"
    assert h._materialsOnItemTab == ["9-F1", "9-F2"]


def test_write_data_9f_without_tab_is_swallowed(tmp_path):
    h = make(tmp_path)
    p = tmp_path / "w.txt"
    with h.open_writer(str(p)) as f:
        h.write_data(f, "9-Fnotab", False)  # Substring(0,-1) throws, caught (780)
    assert len(h.errors) == 1 and "LegacyRangeError" in h.errors[0]
    assert read(p) == ""
    assert issubclass(LegacyRangeError, Exception)


def test_log_message_is_noop():
    assert IoMixin.log_message("x") is None


# --- price formula / product code (986-1088) -----------------------------------------------------
def test_write_price_formula_data(tmp_path):
    rows = [{"Site": "UK", "CustPriceCode": "C1", "ItemPriceCode": "I'1", "PriceFormula": "F", "FirstPrice": 1.5,
             "EffDate": "01 Jan 2020"}]
    h = make(tmp_path, {"FROM PriceFormula pf": rows}, _siteId=3, _priceCodes=["I'1", "I2"],
             _priceformulaFile=str(tmp_path / "pf.txt"))
    h.write_price_formula_data()
    assert read(tmp_path / "pf.txt") == (
        "Site\tCustPriceCode\tItemPriceCode\tPriceFormula\tPriceFactor\tEffectiveDate\r\n"
        "UK\tC1\tI'1\tF\t1.5\t01 Jan 2020\r\n")
    conn, sql, params, stage = h.db.calls[0]
    assert params == (3, "I'1", "I2")
    assert "I'1" not in sql and sql.count("pm.ItemPriceCode = ?") == 2 and " OR " in sql
    assert sql.endswith(") ORDER BY pm.CustPriceCode, pf.EffectiveDate DESC")


def test_write_price_formula_no_codes_header_only_no_query(tmp_path):
    h = make(tmp_path, _priceformulaFile=str(tmp_path / "pf.txt"))
    h.write_price_formula_data()
    assert read(tmp_path / "pf.txt").count("\r\n") == 1
    assert h.db.calls == []


def test_write_product_code_data_header_has_five_columns_rows_four(tmp_path):
    rows = [{"Site": "UK", "Product_Code": "P1", "Description": None, "PriceCode": "PC"}]
    h = make(tmp_path, {"FROM Product_Code pc": rows}, _siteId=2, _productCodes=["P1"],
             _productcodeFile=str(tmp_path / "pc.txt"))
    h.write_product_code_data()
    header, row, _ = read(tmp_path / "pc.txt").split("\r\n")
    assert header.split("\t") == ["Site", "ProductCode", "Description", "PriceCode", "UnitCode"]  # 1051
    assert row.split("\t") == ["UK", "P1", "", "PC"]  # 1069: legacy defect, 4 values
    assert h.db.calls[0][2] == (2, "P1") and "P1" not in h.db.calls[0][1]


def test_write_product_code_db_error_swallowed_and_file_closed(tmp_path):
    def boom(conn, sql, params):
        raise RuntimeError("db down")

    h = make(tmp_path, {"FROM Product_Code pc": boom}, _productCodes=["P"],
             _productcodeFile=str(tmp_path / "pc.txt"))
    h.write_product_code_data()
    assert len(h.errors) == 1
    assert read(tmp_path / "pc.txt").startswith("Site\tProductCode")


def test_write_product_code_unwritable_path_propagates(tmp_path):
    h = make(tmp_path, _productcodeFile=str(tmp_path / "nodir" / "pc.txt"))
    with pytest.raises(OSError):  # StreamWriter opened outside the try (1041)
        h.write_product_code_data()


# --- ConvertToXLS (523-707) ----------------------------------------------------------------------
def _files(h, tmp_path, with_all=True):
    h._exportFileName = "Exp"
    h.allocate_files()
    w = lambda p, s: open(p, "w", encoding="utf-8", newline="").write(s)
    w(h._itemFile, "Item\tDesc\r\nA\tx\r\n")
    w(h._materialFile, "Parent\tFeat\tComp\tOpt\r\nA\t1\tC\t007\r\n")
    if with_all:
        w(h._superprodFile, "Item\tSeq\r\nS\t1\r\n")
        w(h._planningFile, "Item\tDesc\r\n")
        w(h._descriptionFile, "Lang\tItem\tDescription\r\nen\tA\td\r\n")
        w(h._productcodeFile, "Site\tProductCode\r\nUK\tP\r\n")
        w(h._priceformulaFile, "Site\tCust\r\nUK\tC\r\n")


def test_xlsx_syteline_sheets_and_materials_text_forcing(tmp_path):
    h = make(tmp_path, _syteline=True)
    _files(h, tmp_path)
    h.convert_to_xls()
    assert h.errors == [] and h.warnings == []
    wb = load_workbook(tmp_path / "Exp.xlsx")
    assert wb.sheetnames == ["Items", "Materials", "SPextract", "Planning Data", "Descriptions"]
    assert wb.active.title == "Items"
    mat = wb["Materials"]
    assert mat["D2"].value.rstrip("\r\n") == "007" and mat["D2"].data_type == "s" and mat["D2"].quotePrefix  # 506-508
    assert mat["D1"].quotePrefix and not mat["C2"].quotePrefix
    # Items from Split('\n'): trailing CR stays on last cell, extra empty last row (579-580)
    items = wb["Items"]
    assert items["B1"].value.rstrip("\r\n") == "Desc" and items["B1"].value != "Desc" and items.max_row == 3
    # SPextract from ReadAllLines: no CR, no phantom row
    assert wb["SPextract"]["B2"].value == "1" and wb["SPextract"].max_row == 2
    assert wb["Descriptions"]["C2"].value == "d"


def test_xlsx_syteline_csi_only_items_and_materials(tmp_path):
    h = make(tmp_path, _syteline=True, _formatCSI=True)
    _files(h, tmp_path, with_all=False)
    h.convert_to_xls()
    wb = load_workbook(tmp_path / "Exp.xlsx")
    assert wb.sheetnames == ["Items", "Materials"]  # 548: CSI skips the three extra sheets


def test_xlsx_non_syteline_sheets(tmp_path):
    h = make(tmp_path, _syteline=False)
    _files(h, tmp_path)
    h.convert_to_xls()
    wb = load_workbook(tmp_path / "Exp.xlsx")
    assert wb.sheetnames == ["Items", "Materials", "PriceFormula", "ProductCode", "SPextract"]


def test_xlsx_replaces_previous_file_and_truncates_long_rows(tmp_path):
    h = make(tmp_path, _syteline=True, _formatCSI=True)
    _files(h, tmp_path, with_all=False)
    (tmp_path / "Exp.xlsx").write_bytes(b"old")
    open(h._itemFile, "w", newline="").write("H1\tH2\r\nA\tB\tEXTRA\r\nC\r\n")
    h.convert_to_xls()
    items = load_workbook(tmp_path / "Exp.xlsx")["Items"]
    assert items.max_column == 2  # 495 / 503: width from first line, longer rows truncated
    assert items["A3"].value.rstrip("\r\n") == "C" and items["B3"].value is None


def test_xlsx_failure_recorded_and_warned(tmp_path):
    h = make(tmp_path, _syteline=False)
    _files(h, tmp_path)
    import os
    os.remove(h._superprodFile)  # ReadAllLines throws (566) -> MsgBox, run continues (615-621)
    h.convert_to_xls()
    assert len(h.errors) == 1 and "ERROR during XLS export" in h.errors[0]
    assert len(h.warnings) == 1 and "failed" in h.warnings[0]
    assert not (tmp_path / "Exp.xlsx").exists()


def test_write_to_sheet_empty_lines_and_slow_variant(tmp_path):
    from openpyxl import Workbook
    h = make(tmp_path)
    ws = Workbook().active
    h.write_to_sheet(ws, [], "Items")  # 490: returns
    assert ws.max_row == 1 and ws["A1"].value is None
    msgs = []
    h._status_callback = msgs.append
    h.write_to_sheet_slow(ws, ["a\tb\r", "c"], "ctx")  # 476: CR removed
    assert ws["B1"].value == "b" and ws["A2"].value == "c"
    assert msgs[0] == "writing row 1 of 2 (ctx)"
